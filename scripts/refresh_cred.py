"""TeamLink credentialing refresh.

Usage:
  python3 scripts/refresh_cred.py --ahad "Credentialing Status by Ahad.xlsx" --abrar "Abrar and Sarmad Practices.xlsx"

Sources (Google Drive):
  Ahad   : https://docs.google.com/spreadsheets/d/1-6S_02qhI2lJb_n4dSgIJoWHOpc7wv2w  (xlsx owned by Abdul.Ahad@teletherapeuticsinc.com)
  Abrar  : https://docs.google.com/spreadsheets/d/10PNxwrUoviH-OhNoV3nVf6Ro3rOtBxMMTd29BrunY3Y  (Google Sheet owned by SarmadAli.Khan@teletherapeuticsinc.com; export as xlsx)

Status rules: Ahad tabs have an explicit group Status column (taken as-is; a payer is In-Network if the group is, or if
>=2 providers are individually In-Network). Abrar/Sarmad tabs are narrative cells, keyword-classified on the most recent
dated note; best status across the group cell and provider cells wins. Entity/footer rows are skipped.
"""
import openpyxl, re, json, sys, collections

VALID = ['In-Network', 'In-Process', 'Pending', 'REQ Denied', 'Out of Network']
RANK = {s: i for i, s in enumerate(VALID)}  # lower = better

def clean(c):
    if c is None: return ''
    return re.sub(r'\s+', ' ', str(c).replace('\xa0', ' ')).strip()

HEADER_WORDS = {'insurances', 'insurance', 'payers name', 'payer name', 'payers', 'payer'}
JUNK_PAYER = re.compile(r'^(group name|tax id|npi|address|s\. no\.?|no\. of providers|license|last caqh|caqh|total|notes?|comments?|practice|mailing)', re.I)
JUNK_FULL = {'therapycentral inc', 'well-america inc', 'teletherapeutics health inc', 'group', 'group contract', 'group contracts'}

# ---- narrative classifier (Abrar/Sarmad format) ----
DATE_RE = re.compile(r'(\d{1,2}/\d{1,2}/\d{2,4})\s*[:\-–]?')

def latest_segment(text):
    """Return the text segment that follows the most recent date stamp, else whole text."""
    parts = [(m.start(), m.group(1)) for m in DATE_RE.finditer(text)]
    if len(parts) < 2: return text
    def key(d):
        m, d_, y = d.split('/'); y = int(y); y = y + 2000 if y < 100 else y
        return (y, int(m), int(d_))
    best = max(parts, key=lambda p: key(p[1]))
    return text[best[0]:]

RULES = [
    ('Out of Network', [r'out[\s\-]*of[\s\-]*network', r'\boon\b', r'non[\s\-]*par\b', r'opt(ed)?[\s\-]*out', r'not a participating']),
    ('REQ Denied',     [r'denied', r'declin', r'reject', r'not accepting', r'panel (is )?closed', r'closed (panel|network|to new)', r'network is closed',
                        r'no credential(l)?ing needed', r'not credential', r'does not credential', r'do not credential', r'don\'t credential',
                        r'(require|requires|required|need|needs|needed|at ?least)\D{0,25}\d+\s*(providers?|clinicians?)', r'\d+\s*providers?\s*(are|is)?\s*(required|needed)',
                        r'third[\s\-]*party administrator', r'\btpa\b', r'not (in|under) (our|the) (contract|network)', r'terminated', r'termination', r'unable to (credential|contract)',
                        r'no longer', r'not available', r'not offering', r'cannot (be )?credential', r'can\'?t (be )?credential', r'not required', r'no need', r'pharmacy plan', r'not op(e)?rating', r'does not operate', r'not (in|serving) (our|the|this) (area|county|state)']),
    ('In-Network',     [r'in[\s\-]*network', r'innetwork', r'\bapproved\b', r'\bapprove\b', r'got approved', r'contracted', r'contract (has been|is|was) (received|executed|signed|finalized|approved)',
                        r'participating', r'credentialed', r'\beff(ective)?\.?\s*(date)?\s*:?\s*\d', r'effective (date|from)', r'provider id\s*:?\s*[a-z0-9]{4,}', r'\bpar\b', r'accepted', r'loaded', r'welcome letter', r'active']),
    ('In-Process',     [r'in[\s\-]*process', r'in progress', r'processing', r'submitted', r'submission', r'application', r'request (to join|has been|was|id)', r'under review', r'reviewing',
                        r'sent (to|the|an?)', r'applied', r'request id', r'case (id|#|number)', r'reference', r'ref\s*:', r'ticket', r'initiated', r'started', r'recredential']),
    ('Pending',        [r'pending', r'on hold', r'\bhold\b', r'waiting', r'await', r'need', r'require', r'missing', r'follow[\s\-]*up', r'complaint', r'not responding', r'no response',
                        r'reach(ed)? out', r'called', r'email(ed)?', r'escalat', r'expired', r'login', r'caqh', r'dea', r'\bcoi\b', r'will (be )?(submit|apply|start)', r'to be (submitted|applied)', r'tbd', r'\?']),
]
COMPILED = [(s, [re.compile(p, re.I) for p in pats]) for s, pats in RULES]

def classify_text(t):
    t = clean(t)
    if not t: return None
    low = t.lower()
    for s in VALID:
        if low == s.lower() or low == s.lower().replace('-', ' '): return s
    seg = latest_segment(t)
    for text in (seg, t):
        for status, pats in COMPILED:
            if any(p.search(text) for p in pats):
                return status
    return None

def best(statuses):
    statuses = [s for s in statuses if s]
    if not statuses: return None
    return min(statuses, key=lambda s: RANK[s])

ENTITY_RE = re.compile(r'\b(inc\.?|corp(oration)?\.?|llc|l\.l\.c\.|group npi|tax id)\b', re.I)
IDLIKE_RE = re.compile(r'^(group npi:?\s*)?\d{2}-?\d{7}(\.0)?$|^\d{9,10}(\.0)?$')
def is_entity_row(payer, cells):
    """Footer rows listing the practice entity (NPI / tax id / address) rather than a payer."""
    nonempty = [c for c in cells if c]
    if ENTITY_RE.search(payer) and (not nonempty or all(IDLIKE_RE.match(c) or classify_text(c) is None for c in nonempty)):
        return True
    if nonempty and all(IDLIKE_RE.match(c) for c in nonempty):
        return True
    return False

# ---- Ahad format ----
AHAD_BRAND = {'tele': 'Teletherapeutics', 'well a': 'Well-America', 'therapy': 'Therapy Central'}

def parse_ahad(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []; unclassified = []
    for ws in wb.worksheets:
        title = ws.title.strip()
        m = re.match(r'^(Tele|Well A|THERAPY)\s*-\s*([A-Z]{2})$', title, re.I)
        if not m: continue
        brand = AHAD_BRAND[m.group(1).lower()]; state = m.group(2).upper()
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = next((i for i, r in enumerate(rows) if r and clean(r[0]).lower() in HEADER_WORDS), None)
        if hdr is None:
            print('!! no header', title); continue
        header = [clean(c).lower() for c in rows[hdr]]
        status_cols = [i for i, h in enumerate(header) if h == 'status']
        group_col = status_cols[0] if status_cols else 2
        seen = set()
        for r in rows[hdr + 1:]:
            if not r: continue
            payer = clean(r[0])
            if not payer or JUNK_PAYER.match(payer) or payer.lower() in JUNK_FULL: continue
            gs_raw = clean(r[group_col]) if group_col < len(r) else ''
            if is_entity_row(payer, [clean(r[i]) for i in range(1, min(len(r), 8))]): continue
            # Group status column is the team's call; but if the group isn't contracted and >=2 providers are
            # individually In-Network, the practice can still bill that payer -> In-Network (one stray cell is ignored).
            g = classify_text(gs_raw)
            prov = [classify_text(r[i]) for i in status_cols[1:] if i < len(r)]
            inn = sum(1 for x in prov if x == 'In-Network')
            gs = 'In-Network' if (g == 'In-Network' or inn >= 2) else (g or best(prov))
            if gs is None:
                if gs_raw: unclassified.append((title, payer, gs_raw[:80]))
                gs = 'Pending'  # listed, no status recorded anywhere
            key = payer.lower()
            if key in seen: continue
            seen.add(key)
            out.append({'brand': brand, 'state': state, 'payer': payer, 'status': gs})
    return out, unclassified

# ---- Abrar/Sarmad format ----
SKIP_ABRAR = {'practices and providers', 'rates', 'blockers & gaps', 'renewal watch (90 days)'}
STATE_NAME = {'california': 'CA'}

def parse_abrar(path, brand='Well-America'):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []; unclassified = []
    for ws in wb.worksheets:
        title = ws.title.strip()
        if title.lower() in SKIP_ABRAR: continue
        state = STATE_NAME.get(title.lower(), title.upper())
        if not re.match(r'^[A-Z]{2}$', state): print('!! skip tab', title); continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = next((i for i, r in enumerate(rows) if r and clean(r[0]).lower() in HEADER_WORDS), None)
        if hdr is None:
            print('!! no header', title); continue
        header = [clean(c).lower() for c in rows[hdr]]
        group_col = next((i for i, h in enumerate(header) if h.startswith('group')), None)
        prov_cols = [i for i, h in enumerate(header) if i > 0 and i != group_col and h]
        seen = set()
        for r in rows[hdr + 1:]:
            if not r: continue
            payer = clean(r[0])
            if not payer or JUNK_PAYER.match(payer) or payer.lower() in JUNK_FULL: continue
            cells = [clean(r[i]) for i in range(1, len(r))]
            if not any(cells): continue  # sub-header / blank row
            gs_raw = clean(r[group_col]) if group_col is not None and group_col < len(r) else ''
            if is_entity_row(payer, cells): continue
            gs = best([classify_text(gs_raw)] + [classify_text(r[i]) for i in prov_cols if i < len(r)])
            if gs is None:
                unclassified.append((title, payer, (gs_raw or ' | '.join(c for c in cells if c))[:80]))
                gs = 'Pending'
            key = payer.lower()
            if key in seen: continue
            seen.add(key)
            out.append({'brand': brand, 'state': state, 'payer': payer, 'status': gs})
    return out, unclassified


def inject(index_path, data, today):
    """Replace the CRED_DATA block in index.html and update the record counts, date stamp and version tag."""
    data = sorted(data, key=lambda r: (r['state'], r['brand'], r['payer'].lower()))
    s = open(index_path, encoding='utf-8').read()
    start = s.index('// ======= CRED DATA'); i = s.index('const CRED_DATA = [', start); j = s.index('\n];', i) + 3
    old = json.loads(s[i + len('const CRED_DATA = '):j - 1])
    lines = ',\n'.join(json.dumps(r, ensure_ascii=False) for r in data)
    nstates = len({r['state'] for r in data})
    block = (f"// ======= CRED DATA — refreshed {today} by scripts/refresh_cred.py from the Abrar/Sarmad Practices sheet "
             f"and Ahad's Credentialing Status workbook — {len(data)} records, {nstates} states =======\nconst CRED_DATA = [\n{lines}\n];")
    s = s[:start] + block + s[j:]
    oldn, newn = f"{len(old):,}", f"{len(data):,}"
    s = re.sub(r'<div class="home-stat"><div class="n">[\d,]+</div><div class="l">Payer Records</div></div>',
               f'<div class="home-stat"><div class="n">{newn}</div><div class="l">Payer Records</div></div>', s)
    s = re.sub(r'across \d+ states and 3 brands\. Search, filter, and track [\d,]+ payer records',
               f'across {nstates} states and 3 brands. Search, filter, and track {newn} payer records', s)
    import datetime
    pretty = datetime.date.fromisoformat(today).strftime('%B %d, %Y').replace(' 0', ' ')
    s = re.sub(r'<div class="hero-subtitle">[\d,]+ payer records · \d+ states · 3 brands · Updated [^<]*</div>',
               f'<div class="hero-subtitle">{newn} payer records · {nstates} states · 3 brands · Updated {pretty}</div>', s)
    s = re.sub(r'<!-- TeamLink v[\d.]+ [^>]*-->', f'<!-- TeamLink CRED-REFRESH {today} -->', s, count=1)
    open(index_path, 'w', encoding='utf-8').write(s)
    return old

def diff_report(old, new):
    o = {(r['brand'], r['state'], r['payer'].lower()): r['status'] for r in old}
    n = {(r['brand'], r['state'], r['payer'].lower()): r['status'] for r in new}
    trans = collections.Counter((o[k], v) for k, v in n.items() if k in o and o[k] != v)
    lines = [f"Records: {len(old):,} -> {len(new):,} | added {len(set(n) - set(o))} | removed {len(set(o) - set(n))} | status changes {sum(trans.values())}",
             "By status: " + ', '.join(f"{k} {v}" for k, v in collections.Counter(r['status'] for r in new).most_common())]
    for (a, b), c in trans.most_common(): lines.append(f"  {a} -> {b}: {c}")
    added = sorted(set(n) - set(o)); removed = sorted(set(o) - set(n))
    if added: lines.append("Added: " + '; '.join(f"{k[1]} {k[2]}" for k in added))
    if removed: lines.append("Removed: " + '; '.join(f"{k[1]} {k[2]}" for k in removed))
    return '\n'.join(lines)

if __name__ == '__main__':
    import argparse, datetime
    ap = argparse.ArgumentParser(description='Rebuild the TeamLink credentialing dashboard data from the two team workbooks.')
    ap.add_argument('--ahad', required=True, help='Credentialing Status by Ahad.xlsx')
    ap.add_argument('--abrar', required=True, help='Abrar and Sarmad Practices.xlsx')
    ap.add_argument('--index', default='public/index.html')
    ap.add_argument('--today', default=datetime.date.today().isoformat())
    ap.add_argument('--dry-run', action='store_true', help='parse and report, do not write index.html')
    a = ap.parse_args()
    ahad, u1 = parse_ahad(a.ahad)
    abrar, u2 = parse_abrar(a.abrar)
    seen = set(); data = []
    for r in ahad + abrar:
        k = (r['brand'], r['state'], r['payer'].lower())
        if k in seen: continue
        seen.add(k); data.append(r)
    print(f'Parsed: Ahad {len(ahad)}, Abrar/Sarmad {len(abrar)}, merged {len(data)} across {len({r["state"] for r in data})} states')
    if u1 + u2:
        print(f'Unclassified (defaulted to Pending): {len(u1) + len(u2)}')
        for u in (u1 + u2)[:40]: print('  ', u)
    if len(data) < 1000 or len({r['state'] for r in data}) < 20:
        sys.exit('ABORT: parsed data looks truncated — check the source workbooks before overwriting the dashboard.')
    if a.dry_run:
        s = open(a.index, encoding='utf-8').read(); i = s.index('const CRED_DATA = ['); j = s.index('\n];', i)
        old = json.loads(s[i + len('const CRED_DATA = '):j + 2])
    else:
        old = inject(a.index, data, a.today)
    print(diff_report(old, data))
